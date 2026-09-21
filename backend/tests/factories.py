import io

import docx


def make_cv_docx(
    full_name: str,
    email: str,
    phone: str,
    location: str = "Ho Chi Minh City",
    skills: str = "Python, FastAPI, LangChain, RAG, LLM, PostgreSQL, Docker, AWS, Qdrant",
    title: str = "Senior Backend Engineer",
    company: str = "Tech Corp",
    certifications: str | None = None,
    languages: str | None = None,
) -> bytes:
    doc = docx.Document()
    doc.add_paragraph(full_name)
    doc.add_paragraph(email)
    doc.add_paragraph(phone)
    doc.add_paragraph(location)
    doc.add_paragraph("")
    doc.add_paragraph("SUMMARY")
    doc.add_paragraph(f"{title} with 6 years of experience specializing in {skills}.")
    doc.add_paragraph("")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph(skills)
    doc.add_paragraph("")
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph(f"{title} at {company} (Jan 2021 - Present)")
    doc.add_paragraph(f"Delivered projects applying {skills} in a production environment.")
    doc.add_paragraph("")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Bachelor of Science, Ho Chi Minh City University of Technology (2014 - 2018)")
    if certifications:
        doc.add_paragraph("")
        doc.add_paragraph("CERTIFICATIONS")
        doc.add_paragraph(certifications)
    if languages:
        doc.add_paragraph("")
        doc.add_paragraph("LANGUAGES")
        doc.add_paragraph(languages)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def make_cv_with_oversized_certification_line(full_name: str, email: str, phone: str) -> bytes:
    """A CV whose CERTIFICATIONS section contains one very long line — this
    reproduces a real bug where a misclassified/free-text section produced a
    400+ char "certification name" against a VARCHAR(255) column, crashing
    the whole upload with a Postgres StringDataRightTruncation error."""
    doc = docx.Document()
    doc.add_paragraph(full_name)
    doc.add_paragraph(email)
    doc.add_paragraph(phone)
    doc.add_paragraph("")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, FastAPI")
    doc.add_paragraph("")
    doc.add_paragraph("CERTIFICATIONS")
    doc.add_paragraph(
        "Coursera: Machine Learning Operations Specialization, Machine Learning Specialization, "
        "Google Data Analytics Professional, Reinforcement Learning Specialization, Introduction to "
        "Machine Learning on AWS, DeepLearning.AI TensorFlow Developer Professional, Natural Language "
        "Processing Specialization, Deep Learning Specialization, Data Engineering Professional, "
        "LangChain Chat with Your Data — a very long free-text block that easily exceeds two hundred "
        "and fifty five characters, exactly like a real misclassified resume section did in production."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
