import io
import logging

import docx
from pypdf import PdfReader

from app.core.config import settings

logger = logging.getLogger(__name__)

IMAGE_MIME_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}


def extract_text(file_bytes: bytes, extension: str) -> tuple[str, str]:
    """Returns (text, extraction_method) where extraction_method is
    "text" (a normal text layer) or "ocr" (scanned page(s) run through OCR)."""
    extension = extension.lower()
    if extension == ".pdf":
        return _extract_pdf(file_bytes)
    if extension == ".docx":
        return _extract_docx(file_bytes), "text"
    if extension in IMAGE_MIME_TYPES:
        return _extract_image(file_bytes, IMAGE_MIME_TYPES[extension]), "ocr"
    raise ValueError(f"Unsupported extension for text extraction: {extension}")


def _extract_pdf(file_bytes: bytes) -> tuple[str, str]:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages)
    except Exception:
        logger.error("Failed to extract text layer from PDF, will try OCR", exc_info=True)
        text = ""

    if len(text.strip()) >= settings.OCR_MIN_TEXT_LENGTH:
        return text, "text"

    # No (or too little) extractable text layer — this is very likely a
    # scanned CV saved as PDF. Render each page to an image and OCR it.
    logger.info("PDF has little/no text layer (%d chars) — falling back to OCR", len(text.strip()))
    ocr_text = _ocr_pdf(file_bytes)
    if len(ocr_text.strip()) > len(text.strip()):
        return ocr_text, "ocr"
    return text, "text"


def _ocr_pdf(file_bytes: bytes) -> str:
    from pdf2image import convert_from_bytes

    from app.services.ocr.factory import get_ocr_service

    try:
        images = convert_from_bytes(file_bytes, dpi=200)
    except Exception:
        logger.error("Failed to rasterize PDF for OCR", exc_info=True)
        return ""

    ocr_service = get_ocr_service()
    parts = []
    for image in images[: settings.OCR_MAX_PAGES]:
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        parts.append(ocr_service.extract_text(buf.getvalue(), "image/png"))
    return "\n".join(parts)


def _extract_docx(file_bytes: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception:
        logger.error("Failed to extract text from DOCX", exc_info=True)
        raise


def _extract_image(file_bytes: bytes, mime_type: str) -> str:
    from app.services.ocr.factory import get_ocr_service

    return get_ocr_service().extract_text(file_bytes, mime_type)
